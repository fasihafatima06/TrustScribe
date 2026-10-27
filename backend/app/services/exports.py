import io
import json
from datetime import datetime
from typing import List
from ..models import TranscriptSegment, Recording, GlossaryItem

# DOCX generator
def generate_docx_export(recording: Recording, segments: List[TranscriptSegment], glossary_items: List[GlossaryItem] = None) -> io.BytesIO:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()

    # Title
    title = doc.add_heading("TrustScribe Research Verification Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("AI-Assisted Multilingual Transcript Verification & Audit Record")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Metadata Block
    meta_p = doc.add_paragraph()
    meta_p.add_run(f"Project / File: ").bold = True
    meta_p.add_run(f"{recording.filename}\n")
    meta_p.add_run(f"Languages: ").bold = True
    meta_p.add_run(f"{recording.source_language} → {recording.target_language}\n")
    meta_p.add_run(f"Audio Duration: ").bold = True
    meta_p.add_run(f"{int(recording.duration_seconds)} seconds\n")
    meta_p.add_run(f"Generated On: ").bold = True
    meta_p.add_run(f"{datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}\n")

    # Metrics Summary
    total = len(segments)
    reviewed = sum(1 for s in segments if s.review_status == "human_reviewed")
    flagged = sum(1 for s in segments if s.review_status == "review_suggested" or s.is_flagged)
    unclear = sum(1 for s in segments if s.review_status == "unclear")
    unverified = sum(1 for s in segments if s.review_status == "unverified")

    doc.add_heading("1. Verification Summary & Audit Metrics", level=1)
    stats_table = doc.add_table(rows=2, cols=4)
    stats_table.style = 'Table Grid'
    headers = ["Total Segments", "Human Reviewed", "Review Suggested", "Unverified"]
    counts = [str(total), f"{reviewed} ({int(reviewed/total*100) if total else 0}%)", str(flagged), str(unverified)]
    
    for i, h in enumerate(headers):
        stats_table.cell(0, i).paragraphs[0].text = h
        stats_table.cell(0, i).paragraphs[0].runs[0].bold = True
        stats_table.cell(1, i).paragraphs[0].text = counts[i]

    doc.add_paragraph()

    # Executive Summary & Insights
    if recording.summary:
        doc.add_heading("2. Executive Research Summary", level=1)
        doc.add_paragraph(recording.summary)

    if recording.insights:
        doc.add_heading("3. Evidence-Linked Thematic Insights", level=1)
        for ins in recording.insights:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"[{ins.category}] {ins.title}: ").bold = True
            p.add_run(ins.content)
            if ins.linked_timestamp:
                p.add_run(f" (Timestamp: {int(ins.linked_timestamp)}s)").italic = True

    # Glossary Section if present
    if glossary_items and len(glossary_items) > 0:
        doc.add_heading("4. Reusable Project Glossary", level=1)
        g_table = doc.add_table(rows=1, cols=3)
        g_table.style = 'Table Grid'
        g_headers = ["Approved Term", "Category", "Notes"]
        for i, gh in enumerate(g_headers):
            g_table.cell(0, i).paragraphs[0].text = gh
            g_table.cell(0, i).paragraphs[0].runs[0].bold = True
        for g in glossary_items:
            row = g_table.add_row()
            row.cells[0].paragraphs[0].text = g.term
            row.cells[1].paragraphs[0].text = g.category or "General"
            row.cells[2].paragraphs[0].text = g.description or ""
        doc.add_paragraph()

    # Line by Line Verification Record
    doc.add_heading("5. Transcript & Human Review Record", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = 'Table Grid'
    col_names = ["Time", "Speaker", "Language", "Original / Corrected Text", "English Translation", "Review Status"]
    for i, name in enumerate(col_names):
        table.cell(0, i).paragraphs[0].text = name
        table.cell(0, i).paragraphs[0].runs[0].bold = True

    for s in segments:
        row = table.add_row()
        start_m, start_s = int(s.start_time // 60), int(s.start_time % 60)
        end_m, end_s = int(s.end_time // 60), int(s.end_time % 60)
        row.cells[0].paragraphs[0].text = f"{start_m:02d}:{start_s:02d}-{end_m:02d}:{end_s:02d}"
        row.cells[1].paragraphs[0].text = s.speaker_tag or "Speaker"
        row.cells[2].paragraphs[0].text = s.detected_language or "Mixed"
        
        orig_p = row.cells[3].paragraphs[0]
        orig_p.text = s.original_text
        if s.user_edited:
            edit_run = orig_p.add_run("\n[Human Corrected]")
            edit_run.font.size = Pt(8)
            edit_run.font.color.rgb = RGBColor(16, 120, 60)

        row.cells[4].paragraphs[0].text = s.translated_text or ""
        
        status_p = row.cells[5].paragraphs[0]
        status_label = s.review_status.replace("_", " ").title()
        status_p.text = status_label
        if s.flag_reason:
            reason_run = status_p.add_run(f"\nNote: {s.flag_reason}")
            reason_run.font.size = Pt(8)
            reason_run.font.color.rgb = RGBColor(180, 80, 20)

    target = io.BytesIO()
    doc.save(target)
    target.seek(0)
    return target

# PDF generator
def generate_pdf_export(recording: Recording, segments: List[TranscriptSegment]) -> io.BytesIO:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=18, leading=22, textColor=colors.HexColor('#0f172a'))
    sub_style = ParagraphStyle('SubStyle', parent=styles['Normal'], fontSize=9, textColor=colors.HexColor('#64748b'))
    cell_style = ParagraphStyle('CellStyle', parent=styles['Normal'], fontSize=8, leading=10, textColor=colors.HexColor('#1e293b'))
    header_style = ParagraphStyle('HStyle', parent=styles['Normal'], fontSize=8, leading=10, textColor=colors.white, fontName='Helvetica-Bold')
    flag_style = ParagraphStyle('FlagStyle', parent=styles['Normal'], fontSize=7, leading=9, textColor=colors.HexColor('#b45309'))

    elements = []

    # Title & Subtitle
    elements.append(Paragraph("TrustScribe — Multilingual Research Verification Report", title_style))
    elements.append(Paragraph(f"Recording: <b>{recording.filename}</b> | Duration: {int(recording.duration_seconds)}s | Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}", sub_style))
    elements.append(Spacer(1, 12))

    # Executive Summary
    if recording.summary:
        elements.append(Paragraph("<b>Executive Summary:</b> " + recording.summary, cell_style))
        elements.append(Spacer(1, 10))

    # Stats block
    total = len(segments)
    reviewed = sum(1 for s in segments if s.review_status == "human_reviewed")
    flagged = sum(1 for s in segments if s.review_status == "review_suggested" or s.is_flagged)
    
    stats_data = [
        ["Total Segments", "Human Reviewed", "Review Suggested", "Verification Progress"],
        [str(total), f"{reviewed} lines", f"{flagged} lines", f"{int(reviewed/total*100) if total else 0}% completed"]
    ]
    t_stats = Table(stats_data, colWidths=[130, 130, 130, 140])
    t_stats.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e293b')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('BACKGROUND', (0,1), (-1,1), colors.HexColor('#f8fafc')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
    ]))
    elements.append(t_stats)
    elements.append(Spacer(1, 14))

    # Segment Table
    table_data = [
        [
            Paragraph("<b>Time</b>", header_style),
            Paragraph("<b>Speaker</b>", header_style),
            Paragraph("<b>Original Text / Audio</b>", header_style),
            Paragraph("<b>English Translation</b>", header_style),
            Paragraph("<b>Verification Status</b>", header_style)
        ]
    ]

    for s in segments:
        start_m, start_s = int(s.start_time // 60), int(s.start_time % 60)
        time_str = f"{start_m:02d}:{start_s:02d}"
        
        orig_content = s.original_text
        if s.user_edited:
            orig_content += " <i>[Corrected]</i>"
        
        status_desc = s.review_status.replace("_", " ").title()
        if s.flag_reason:
            status_desc += f"<br/><font color='#b45309'>Note: {s.flag_reason}</font>"

        table_data.append([
            Paragraph(time_str, cell_style),
            Paragraph(f"<b>{s.speaker_tag}</b><br/><font color='#64748b'>{s.detected_language}</font>", cell_style),
            Paragraph(orig_content, cell_style),
            Paragraph(s.translated_text or "—", cell_style),
            Paragraph(status_desc, cell_style)
        ])

    table = Table(table_data, colWidths=[40, 80, 180, 150, 90])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#2563eb')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer

# TXT generator
def generate_txt_export(recording: Recording, segments: List[TranscriptSegment]) -> str:
    total = len(segments)
    reviewed = sum(1 for s in segments if s.review_status == "human_reviewed")
    flagged = sum(1 for s in segments if s.review_status == "review_suggested" or s.is_flagged)

    lines = [
        "================================================================================",
        "                       TRUSTSCRIBE RESEARCH VERIFICATION REPORT                 ",
        "================================================================================",
        f"Recording:        {recording.filename}",
        f"Languages:        {recording.source_language} -> {recording.target_language}",
        f"Duration:         {int(recording.duration_seconds)} seconds",
        f"Generated:        {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}",
        f"Review Progress:  {reviewed}/{total} segments verified ({int(reviewed/total*100) if total else 0}%)",
        f"Flagged Lines:    {flagged} segments require verification attention",
        "================================================================================",
        ""
    ]

    if recording.summary:
        lines.append("EXECUTIVE SUMMARY:")
        lines.append(recording.summary)
        lines.append("")

    if recording.insights:
        lines.append("KEY THEMATIC INSIGHTS:")
        for ins in recording.insights:
            ts_str = f" [{int(ins.linked_timestamp)}s]" if ins.linked_timestamp else ""
            lines.append(f"- [{ins.category}] {ins.title}{ts_str}: {ins.content}")
        lines.append("")

    lines.append("VERIFIED TRANSCRIPT & TRANSLATION LOG:")
    lines.append("-" * 80)

    for seg in segments:
        s_m, s_s = int(seg.start_time // 60), int(seg.start_time % 60)
        e_m, e_s = int(seg.end_time // 60), int(seg.end_time % 60)
        time_tag = f"[{s_m:02d}:{s_s:02d} - {e_m:02d}:{e_s:02d}]"
        
        status_label = seg.review_status.upper().replace("_", " ")
        lines.append(f"{time_tag} {seg.speaker_tag} ({seg.detected_language}) | STATUS: {status_label}")
        lines.append(f"  ORIGINAL:    {seg.original_text}")
        lines.append(f"  TRANSLATION: {seg.translated_text or 'N/A'}")
        if seg.raw_engine1_text and seg.raw_engine2_text:
            lines.append(f"  [Engine 1: {seg.raw_engine1_text}]")
            lines.append(f"  [Engine 2: {seg.raw_engine2_text}]")
            lines.append(f"  [Match Score: {int(seg.comparison_score * 100)}%]")
        if seg.is_flagged and seg.flag_reason:
            lines.append(f"  [REVIEW WARNING: {seg.flag_reason}]")
        if seg.user_edited:
            lines.append(f"  [HUMAN AUDIT: Edited and verified by {seg.reviewer_id or 'Researcher'}]")
        lines.append("")

    return "\n".join(lines)

# JSON generator
def generate_json_export(recording: Recording, segments: List[TranscriptSegment]) -> str:
    data = {
        "report_type": "TrustScribe Verification Record",
        "exported_at": datetime.utcnow().isoformat(),
        "recording": {
            "id": recording.id,
            "filename": recording.filename,
            "duration_seconds": recording.duration_seconds,
            "source_language": recording.source_language,
            "target_language": recording.target_language,
            "summary": recording.summary
        },
        "insights": [
            {
                "category": i.category,
                "title": i.title,
                "content": i.content,
                "linked_timestamp": i.linked_timestamp
            } for i in (recording.insights or [])
        ],
        "segments": [
            {
                "id": s.id,
                "order": s.segment_order,
                "start_time": s.start_time,
                "end_time": s.end_time,
                "speaker": s.speaker_tag,
                "language": s.detected_language,
                "original_text": s.original_text,
                "translated_text": s.translated_text,
                "engine1_text": s.raw_engine1_text,
                "engine2_text": s.raw_engine2_text,
                "comparison_score": s.comparison_score,
                "confidence_score": s.confidence_score,
                "review_status": s.review_status,
                "is_flagged": s.is_flagged,
                "flag_reason": s.flag_reason,
                "user_edited": s.user_edited,
                "reviewer_id": s.reviewer_id,
                "reviewed_at": s.reviewed_at.isoformat() if s.reviewed_at else None
            } for s in segments
        ]
    }
    return json.dumps(data, indent=2, ensure_ascii=False)
